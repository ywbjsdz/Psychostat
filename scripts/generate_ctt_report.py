"""Create a Chinese results-style CTT report from Psychostat CSV/PNG outputs."""
from __future__ import annotations
import argparse, base64, os
from pathlib import Path
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def read_csv(folder: Path, name: str) -> pd.DataFrame:
    path = folder / name
    return pd.read_csv(path, encoding="utf-8") if path.exists() else pd.DataFrame()

def fmt(v):
    if pd.isna(v): return "—"
    if isinstance(v, float):
        if abs(v) < .001 and v != 0: return "< .001"
        return f"{v:.3f}".lstrip("0") if abs(v) < 1 else f"{v:.3f}"
    return str(v)

def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr(); borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None: borders = OxmlElement("w:tcBorders"); tcPr.append(borders)
    for edge, data in kwargs.items():
        tag = "w:" + edge; element = borders.find(qn(tag))
        if element is None: element = OxmlElement(tag); borders.append(element)
        for key, value in data.items(): element.set(qn("w:" + key), str(value))

def three_line_table(doc: Document, df: pd.DataFrame, title: str, max_rows=30):
    doc.add_paragraph(title, style="Caption")
    if df.empty:
        doc.add_paragraph("未提供或该分析未执行。")
        return
    df = df.head(max_rows).copy(); table = doc.add_table(rows=1, cols=len(df.columns)); table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, col in enumerate(df.columns): hdr[i].text = str(col)
    for _, row in df.iterrows():
        cells = table.add_row().cells
        for i, val in enumerate(row): cells[i].text = fmt(val)
    # Three horizontal rules: top, header-bottom, and bottom only.
    all_rows = table.rows
    for c in all_rows[0].cells:
        set_cell_border(c, top={"val":"single","sz":"10","color":"000000"}, bottom={"val":"single","sz":"6","color":"000000"}, left={"val":"nil"}, right={"val":"nil"})
    for row in all_rows[1:-1]:
        for c in row.cells: set_cell_border(c, left={"val":"nil"}, right={"val":"nil"}, top={"val":"nil"}, bottom={"val":"nil"})
    for c in all_rows[-1].cells:
        set_cell_border(c, bottom={"val":"single","sz":"10","color":"000000"}, left={"val":"nil"}, right={"val":"nil"})
    doc.add_paragraph("")

def heading(doc, text, level=1): doc.add_heading(text, level=level)
def figure(doc, path: Path, caption: str):
    if path.exists():
        doc.add_picture(str(path), width=Inches(6.2)); p=doc.add_paragraph(caption); p.alignment=WD_ALIGN_PARAGRAPH.CENTER

def main():
    ap=argparse.ArgumentParser(); ap.add_argument("--result-dir-utf8-base64", required=True); args=ap.parse_args()
    folder=Path(base64.b64decode(args.result_dir_utf8_base64).decode("utf-8")); doc=Document()
    normal=doc.styles["Normal"]; normal.font.name="Microsoft YaHei"; normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei"); normal.font.size=Pt(10.5)
    title=doc.add_heading("Psychostat 经典测量理论（CTT）分析报告", 0); title.alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("自动生成初稿｜请结合量表理论、内容效度与独立样本复核后用于正式报告。", style="Subtitle")
    clean=read_csv(folder,"01_cleaning_summary.csv"); item=read_csv(folder,"02_item_analysis.csv"); rel=read_csv(folder,"03_reliability.csv"); alpha=read_csv(folder,"03_alpha_if_deleted.csv"); criterion=read_csv(folder,"03_criterion_validity.csv"); known=read_csv(folder,"03_known_group_validity.csv"); diag=read_csv(folder,"04_efa_diagnostics.csv"); loads=read_csv(folder,"04_efa_rotated_loadings.csv"); variance=read_csv(folder,"04_efa_variance.csv"); hist=read_csv(folder,"04_efa_deletion_history.csv"); fit=read_csv(folder,"05_cfa_fit.csv"); cfa_loads=read_csv(folder,"05_cfa_loadings.csv"); conv=read_csv(folder,"05_cfa_convergent_validity.csv"); disc=read_csv(folder,"05_cfa_discriminant_validity.csv")
    heading(doc,"摘要")
    if not clean.empty:
        r=clean.iloc[0]; doc.add_paragraph(f"本分析对 {int(r.raw_n)} 名被试的量表数据进行了预设的 CTT 流程：数据清洗、项目分析、信效度检验、探索性因子分析（EFA）和验证性因子分析（CFA）。清洗后保留 {int(r.retained_n)} 名被试；以下自动输出仅作为量表修订的量化证据，不能替代题目内容与理论判断。")
    heading(doc,"1. 数据清洗与质量控制")
    doc.add_paragraph("规则：个体缺失率 <5% 采用题目中位数插补；>20% 剔除；5%–20% 不自动决定，须在配置中明确选择 listwise 或 median_impute。反向题按 k + 1 − 原始得分换算。注意力题、过短作答时间、直线作答和总分极端值均保留审计记录。")
    three_line_table(doc, clean, "表 1\n数据清洗汇总")
    heading(doc,"2. 项目分析")
    doc.add_paragraph("使用总分前后 27% 极端组进行独立样本 t 检验。CR 决断值主列（CR_t、df、p）采用合并方差 t 检验（df = n高 + n低 − 2，与 SPSS/教材决断值口径一致）；Welch 校正结果另见 CR_t_welch、df_welch 列。建议保留 CR > 3 且 p < .05、并且校正后题总相关（CITC）>.30 的题目；自动建议不应取代内容效度审查。")
    three_line_table(doc, item, "表 2\n项目区分度与 CITC", max_rows=50)
    heading(doc,"3. 信度与外部效度")
    doc.add_paragraph("Cronbach’s α > .70 视为可接受；若删除某题后 α 上升 > .05，工具给出删题建议。alpha_ci_lower/alpha_ci_upper 为 α 的 95% 正态近似置信区间；omega 为单因子 McDonald’s ω（计算失败或题目不足时为 NA）。若提供校标变量，将检验相关方向和显著性；若提供已知组变量，将进行 Welch t 检验或单因素方差分析并报告效应量。")
    three_line_table(doc, rel, "表 3\nCronbach’s α")
    if not rel.empty and "source" in rel.columns and rel["source"].astype(str).str.contains("EFA-derived", na=False).any():
        doc.add_paragraph("注意：分量表维度来自同一批数据的 EFA，其 α 偏乐观，应在独立样本验证后再报告。")
    three_line_table(doc, alpha, "表 4\n删除项目后的 α")
    three_line_table(doc, criterion, "表 5\n校标关联效度")
    three_line_table(doc, known, "表 6\n已知组效度")
    heading(doc,"4. 探索性因子分析（EFA）")
    doc.add_paragraph("EFA 的前提通常为 KMO > .60 且 Bartlett 球形检验 p < .05（本工具在未达标时给出警告并继续分析，提醒解读需谨慎）。因子数默认采用平行分析确定（以碎石图辅助判断，可由配置显式指定）；本流程使用配置的 Varimax 或 Promax 旋转。主载荷<.40、交叉载荷、共同度<.20 或因子题数不足的题目仅作“建议删除”标注，不会自动删除——是否删题由研究者结合理论决定。累计方差应超过 50%。")
    three_line_table(doc, diag, "表 7\nEFA 前提检验")
    three_line_table(doc, variance, "表 8\nEFA 方差解释")
    three_line_table(doc, hist, "表 9\nEFA 建议删除题目标注（仅标注，未自动删除）")
    figure(doc, folder/"efa_scree_plot.png", "图 1. 碎石图。虚线表示特征值 1。")
    figure(doc, folder/"efa_loading_heatmap.png", "图 2. 旋转后因子载荷热图。")
    three_line_table(doc, loads, "表 10\n旋转后因子载荷", max_rows=50)
    heading(doc,"5. 验证性因子分析（CFA）")
    doc.add_paragraph("CFA 判断标准为 χ²/df<5、CFI/TLI>.90、RMSEA<.08、SRMR<.08；聚合效度要求标准化载荷>.50、CR>.70、AVE>.50；区分效度要求因子相关<.85 且 AVE 平方根大于因子间相关。本次若直接以同一样本 EFA 结构进行 CFA，结果仅为探索性内部验证；正式验证应使用独立样本或预先指定模型。")
    three_line_table(doc, fit, "表 11\nCFA 拟合指标")
    three_line_table(doc, cfa_loads, "表 12\nCFA 标准化载荷与显著性（heywood_warning=TRUE 表示该行载荷绝对值>1 或方差估计为负，属不合规解，需谨慎报告）", max_rows=50)
    three_line_table(doc, conv, "表 13\n聚合效度")
    three_line_table(doc, disc, "表 14\n区分效度")
    figure(doc, folder/"cfa_path_diagram.png", "图 3. CFA 标准化路径图。")
    heading(doc,"6. 结果段落（可编辑初稿）")
    if not rel.empty:
        alpha_total=rel.iloc[0].get("alpha", float("nan")); doc.add_paragraph(f"量表总分的内部一致性为 Cronbach’s α = {fmt(alpha_total)}。项目分析按总分前后 27% 的极端组进行（CR 决断值为合并方差 t 检验，df = n高 + n低 − 2），题目保留建议同时依据 CR > 3、p < .05 与 CITC > .30。")
    if not diag.empty:
        d=diag.iloc[0]; doc.add_paragraph(f"EFA 前提检验显示 KMO = {fmt(d.get('KMO'))}，Bartlett 球形检验 χ²({fmt(d.get('Bartlett_df'))}) = {fmt(d.get('Bartlett_chisq'))}, p {fmt(d.get('Bartlett_p'))}。")
    if not fit.empty:
        f=fit.iloc[0]; doc.add_paragraph(f"CFA 的拟合为 χ²({fmt(f.get('df'))}) = {fmt(f.get('chisq'))}，χ²/df = {fmt(f.get('chisq_df'))}，CFI = {fmt(f.get('CFI'))}，TLI = {fmt(f.get('TLI'))}，RMSEA = {fmt(f.get('RMSEA'))}，SRMR = {fmt(f.get('SRMR'))}。")
    heading(doc,"7. 解释限制与建议")
    doc.add_paragraph("自动阈值仅用于初筛。删题前应复核理论涵义、内容覆盖和题目措辞；不要仅因统计阈值而删除核心内容。缺失的机制（MCAR/MAR/MNAR）、样本量、Likert 数据的有序性质及模型识别均会影响结果。若 EFA 和 CFA 使用同一批数据，不能作为严格的验证性证据。")
    doc.add_paragraph("来源说明：本报告的表格、统计量和图形来自 Psychostat 的 R 分析输出；报告语言为自动生成草稿。")
    output=folder/"ctt_report_zh.docx"; doc.save(output); print(f"Created report: {output}")

if __name__ == "__main__": main()
