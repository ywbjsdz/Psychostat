"""Create Chinese and English results-style reports from Psychostat statistics outputs.

Every table produced by a method module is embedded (except raw data files), and each
section carries narrative paragraphs: what the method does, what the data look like,
the results-style result sentence, and a plain-language conclusion.
"""
from __future__ import annotations
import argparse, base64, glob, os, re, sys
from pathlib import Path
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# ── docx helpers (three-line-table style shared with the CTT generator) ──
def read_csv(folder: Path, pattern: str) -> pd.DataFrame:
    hits = sorted(glob.glob(str(folder / pattern)))
    return pd.read_csv(hits[0], encoding="utf-8") if hits else pd.DataFrame()

# table-cell formatting semantics (column names follow the R modules' actual CSV headers):
# integer columns keep integer display; only p-value columns may be abbreviated as "< .001".
INT_COL_NAMES = {"n", "n_total", "n_nonzero_pairs", "count", "freq", "frequency", "observed", "df", "df1", "df2"}
P_COL_RE = re.compile(r"^(?:p|p_.+)$|_p$|_p_|p_?value", re.IGNORECASE)

def is_int_col(col) -> bool:
    """Column with integer semantics (sample size / count / degrees of freedom)."""
    return str(col).strip().lower() in INT_COL_NAMES

def is_p_col(col) -> bool:
    """p-value column: 'p', 'p_...' (p_adjusted), '..._p' (sw_p, levene_p), '_p_' or 'pvalue'.
    Note: pair1/predictor/power etc. start with 'p' but are NOT p-value columns."""
    return bool(P_COL_RE.search(str(col)))

def _is_int_val(fv) -> bool:
    """True when fv is an integer value (tolerance 1e-9); tiny nonzero values that merely
    round to 0 (e.g. p = 8.5e-16) are NOT integers."""
    return fv == 0.0 or (abs(fv - round(fv)) < 1e-9 and round(fv) != 0)

def col_all_int(series) -> bool:
    """Fallback rule: every non-missing value of the column is an integer (|v - round(v)| < 1e-9)."""
    seen = False
    for v in series:
        if pd.isna(v): continue
        seen = True
        try: fv = float(v)
        except (TypeError, ValueError): return False
        if not _is_int_val(fv): return False
    return seen

def fmt(v, integer=False, p_col=False):
    """Format one table cell: integer columns as integers; p columns 结果报告格式 ('< .001', 3
    decimals without leading zero); all other numerics with 3 decimals (leading zero dropped
    when |v| < 1, as before)."""
    if pd.isna(v): return "—"
    if isinstance(v, float):
        if integer and _is_int_val(v): return str(int(round(v)))
        if p_col and abs(v) < .001: return "< .001"
        s = f"{v:.3f}"
        if abs(v) < 1: s = s.lstrip("0")
        return s
    s = str(v)
    try:
        return fmt(float(s), integer=integer, p_col=p_col)
    except ValueError:
        return s

def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr(); borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None: borders = OxmlElement("w:tcBorders"); tcPr.append(borders)
    for edge, data in kwargs.items():
        tag = "w:" + edge; element = borders.find(qn(tag))
        if element is None: element = OxmlElement(tag); borders.append(element)
        for key, value in data.items(): element.set(qn("w:" + key), str(value))

MAX_TABLE_ROWS = 200

def three_line_table(doc: Document, df: pd.DataFrame, title: str, max_rows=MAX_TABLE_ROWS, p_matrix=False):
    """p_matrix=True marks a significance-matrix table (matrix_p_*): its columns carry variable
    names, so every numeric column is a p-value."""
    if title: doc.add_paragraph(title, style="Caption")
    if df.empty:
        doc.add_paragraph("—"); return
    if len(df) > max_rows:
        print(f"[警告] 表格行数超过上限，已截断：{title or '(未命名表)'}｜总行数 {len(df)}｜保留前 {max_rows} 行", file=sys.stderr)
    df = df.head(max_rows).copy(); table = doc.add_table(rows=1, cols=len(df.columns)); table.style = "Table Grid"
    p_flags = [p_matrix or is_p_col(c) for c in df.columns]
    int_flags = [(not pf) and (is_int_col(c) or col_all_int(df[c])) for c, pf in zip(df.columns, p_flags)]
    hdr = table.rows[0].cells
    for i, col in enumerate(df.columns): hdr[i].text = str(col)
    for _, row in df.iterrows():
        cells = table.add_row().cells
        for i, val in enumerate(row): cells[i].text = fmt(val, integer=int_flags[i], p_col=p_flags[i])
    all_rows = table.rows
    for c in all_rows[0].cells:
        set_cell_border(c, top={"val":"single","sz":"10","color":"000000"}, bottom={"val":"single","sz":"6","color":"000000"}, left={"val":"nil"}, right={"val":"nil"})
    for row in all_rows[1:-1]:
        for c in row.cells: set_cell_border(c, left={"val":"nil"}, right={"val":"nil"}, top={"val":"nil"}, bottom={"val":"nil"})
    for c in all_rows[-1].cells:
        set_cell_border(c, bottom={"val":"single","sz":"10","color":"000000"}, left={"val":"nil"}, right={"val":"nil"})
    doc.add_paragraph("")

def figure(doc, path: Path, caption: str):
    if path.exists():
        doc.add_picture(str(path), width=Inches(6.0))
        p = doc.add_paragraph(caption); p.alignment = WD_ALIGN_PARAGRAPH.CENTER

def g1(folder: Path, name: str): return read_csv(folder, f"??_{name}.csv")

def cell(df, row, col, default=None):
    try:
        v = df.iloc[row][col]
        return None if pd.isna(v) else float(v)
    except Exception:
        return default

def num(v, k=2):
    if v is None: return "NA"
    s = f"{v:.{k}f}"
    return s.lstrip("0") if 0 < abs(v) < 1 else s

def result_report(p):
    return "p < .001" if p is not None and p < .001 else ("p = NA" if p is None else "p = " + num(p, 3))

def concl(p, alpha=.05, sig_txt="", ns_txt="", zh=True):
    if p is None: return ""
    t = sig_txt if p < alpha else ns_txt
    return t

ES_D = lambda d: "小" if abs(d) < .2 else ("偏小/接近中等" if abs(d) < .5 else ("中等" if abs(d) < .8 else "大"))
ES_ETA = lambda e: "小" if e < .01 else ("偏小/接近中等" if e < .06 else ("中等" if e < .14 else "大"))
ES_R = lambda r: "小" if abs(r) < .1 else ("偏小/接近中等" if abs(r) < .3 else ("中等" if abs(r) < .5 else "大"))

# ── per-method narrative builders: return (zh_paragraphs, en_paragraphs) ──
def data_line(f, mid, zh=True, var_cols=("mean", "sd"), seg_col="segment"):
    d = g1(f, f"{mid}_desc")
    if d.empty: return None
    if var_cols[0] in d.columns:
        parts = []
        for _, r in d.head(8).iterrows():
            seg = f"（{r[seg_col]}）" if seg_col in d.columns and str(r[seg_col]) != "overall" else ""
            parts.append(f"{r['variable']}{seg}: M = {num(r[var_cols[0]])}, SD = {num(r[var_cols[1]])}" if zh
                         else f"{r['variable']}{seg}: M = {num(r[var_cols[0]])}, SD = {num(r[var_cols[1]])}")
        n = int(d["N"].max()) if "N" in d.columns else None
        lead = f"数据描述：共 {n} 个观测；" if (zh and n) else (f"Data: N = {n}; " if n else ("数据描述：" if zh else "Data: "))
        return lead + ("；".join(parts) if zh else "; ".join(parts)) + ("。" if zh else ".")
    return None

def s_descriptives(f):
    d = g1(f, "descriptives_stats")
    if d.empty: return [], []
    zh = ["本模块对每个变量给出集中量数、离散量数与分布形态，是所有后续分析的第一步。",
          "；".join(f"{r['variable']}：M = {num(r['mean'])}, SD = {num(r['sd'])}，偏度 = {num(r['skewness'])}" for _, r in d.iterrows()) + "。",
          f"结论：|偏度| < 1 的变量视为近似正态；明显偏态的变量（若用于 t/ANOVA）需考虑非参数方法或先做变换。"]
    en = ["This module profiles each variable with central tendency, dispersion and shape.",
          "; ".join(f"{r['variable']}: M = {num(r['mean'])}, SD = {num(r['sd'])}, skewness = {num(r['skewness'])}" for _, r in d.iterrows()) + ".",
          "Conclusion: |skewness| < 1 is treated as approximately normal; clearly skewed outcomes call for nonparametric alternatives or transformations."]
    return zh, en

def s_normality(f):
    d = g1(f, "normality_normality_tests")
    if d.empty: return [], []
    zh = ["本模块检验正态性（与SPSS\"探索\"只放因变量的默认输出一致）。",
          "；".join(f"{r['segment']}：Shapiro-Wilk {result_report(r['sw_p'])}" for _, r in d.iterrows()) + "。（SPSS的K-S显著性显示上限为.200）",
          "结论：" + ("各变量 S-W 均 p ≥ .05，未发现显著偏离正态，参数方法前提可接受。" if all(r['sw_p'] >= .05 for _, r in d.iterrows() if pd.notna(r['sw_p']))
                     else "部分变量/组段 S-W p < .05，谨慎使用参数方法，建议结合直方图、Q-Q图与偏度峰度综合判断。")]
    en = ["This module tests normality, matching SPSS Explore with only the dependent variable.",
          "; ".join(f"{r['segment']}: Shapiro-Wilk {result_report(r['sw_p'])}" for _, r in d.iterrows()) + ".",
          "Conclusion: " + ("No significant departure from normality was detected (all S-W p ≥ .05)."
                            if all(r['sw_p'] >= .05 for _, r in d.iterrows() if pd.notna(r['sw_p']))
                            else "Some variables/segments showed S-W p < .05; check histograms, Q-Q plots and skewness before parametric tests.")]
    has_group = any("=" in str(x) for x in d["segment"])
    zh[0] = "整体正态性检验（不分组）用于判断因变量在全体样本中是否近似正态。" + ("分组检验（按组别）用于判断各水平内是否近似正态；两者均通过时，进一步支持独立样本t检验的正态性前提。" if has_group else "未配置分组检验；独立样本t检验还应结合各组分布与方差齐性判断。")
    zh[2] = "结论：" + ("所有 Shapiro-Wilk 检验 p ≥ .05，未发现显著偏离正态；参数方法前提可接受。" if all(r['sw_p'] >= .05 for _, r in d.iterrows() if pd.notna(r['sw_p'])) else "部分检验 p < .05，需结合直方图、Q-Q图、偏度峰度与研究设计综合判断。")
    en[0] = "Overall normality evaluates the outcome in the full sample." + (" Group-specific tests evaluate normality within each level and jointly support the independent-samples t-test assumption." if has_group else " No group-specific test was requested.")
    en[2] = "Conclusion: " + ("all Shapiro-Wilk tests were non-significant." if all(r['sw_p'] >= .05 for _, r in d.iterrows() if pd.notna(r['sw_p'])) else "some tests were significant; inspect the distribution and design before choosing a parametric test.")
    return zh, en

def s_one_sample_t(f):
    d = g1(f, "one_sample_t_test")
    if d.empty: return [], []
    t, df_, p, dc, md = cell(d, 0, "t"), cell(d, 0, "df"), cell(d, 0, "p"), cell(d, 0, "cohens_d"), cell(d, 0, "mean_difference")
    zh = [f"单样本t检验比较样本均值与检验值：均值差 = {num(md)}，t({int(df_)}) = {num(t)}, {result_report(p)}, Cohen's d = {num(dc)}。",
          concl(p, sig_txt=f"结论：样本均值与检验值差异显著，效应量属于{ES_D(dc)}效应。", ns_txt="结论：无充分证据表明样本均值不同于检验值。")]
    en = [f"The one-sample t test compared the sample mean with the test value: mean difference = {num(md)}, t({int(df_)}) = {num(t)}, {result_report(p)}, Cohen's d = {num(dc)}.",
          concl(p, sig_txt="Conclusion: the sample differs significantly from the test value.", ns_txt="Conclusion: no sufficient evidence of a difference from the test value.")]
    return zh, en

def s_independent_t(f):
    lv = g1(f, "independent_t_levene"); d = g1(f, "independent_t_test"); gs = g1(f, "independent_t_group_stats")
    if d.empty: return [], []
    pl = cell(lv, 0, "levene_p")
    row = 0 if (pl is None or pl >= .05) else 1
    t, df_, p, dc = cell(d, row, "t"), cell(d, row, "df"), cell(d, row, "p"), cell(d, row, "cohens_d")
    grp = "；".join(f"{r['group']}: M = {num(r['mean'])}, SD = {num(r['sd'])}" for _, r in gs.iterrows()) if not gs.empty else ""
    zh = [f"独立样本t检验比较两组均值。{grp}",
          f"Levene检验 {result_report(pl)}，取\"{'假定方差相等' if row == 0 else '不假定方差相等(Welch)'}\"行：t({num(df_, 1)}) = {num(t)}, {result_report(p)}, Cohen's d = {num(dc)}。",
          concl(p, sig_txt=f"结论：两组差异显著，效应量为{ES_D(dc)}效应；报告时写明方向（哪组更高）。", ns_txt="结论：两组均值无显著差异。")]
    en = [f"Independent-samples t test. {grp.replace('；', '; ')}",
          f"Levene's test {result_report(pl)}; using the '{'equal variances assumed' if row == 0 else 'Welch'}' row: t({num(df_, 1)}) = {num(t)}, {result_report(p)}, Cohen's d = {num(dc)}.",
          concl(p, sig_txt="Conclusion: the two groups differ significantly; report the direction.", ns_txt="Conclusion: no significant difference between groups.")]
    return zh, en

def s_paired_t(f):
    d = g1(f, "paired_t_test")
    if d.empty: return [], []
    dlab = str(d.iloc[0].get("difference_direction", "后−前"))
    t, df_, p, dc, md = cell(d, 0, "t"), cell(d, 0, "df"), cell(d, 0, "p"), cell(d, 0, "cohens_d"), cell(d, 0, "mean_difference")
    zh = [f"配对样本t检验比较同一批被试的前后测（差值 = {dlab}）。",
          f"差值均值 = {num(md)}：t({int(df_)}) = {num(t)}, {result_report(p)}, Cohen's d = {num(dc)}。",
          concl(p, sig_txt=f"结论：前后测差异显著（{ES_D(dc)}效应）；注意差值方向为{dlab}，结合变量含义解读升降。", ns_txt="结论：前后测无显著差异。")]
    en = [f"Paired-samples t test (difference = {dlab}).",
          f"Mean difference = {num(md)}, t({int(df_)}) = {num(t)}, {result_report(p)}, Cohen's d = {num(dc)}.",
          concl(p, sig_txt="Conclusion: the pre-post change is significant; interpret the direction accordingly.", ns_txt="Conclusion: no significant pre-post difference.")]
    return zh, en

def s_one_way_anova(f):
    d = g1(f, "one_way_anova_anova"); ph = g1(f, "one_way_anova_posthoc_bonferroni")
    desc = g1(f, "one_way_anova_descriptives")
    iv = str(desc.iloc[0].get("independent_variable", "??")) if not desc.empty else "??"
    dv = str(desc.iloc[0].get("dependent_variable", "???")) if not desc.empty else "???"
    groups = "?".join(f"{r['group']}: M = {num(r['mean'])}, SD = {num(r['sd'])}" for _, r in desc.iterrows()) if not desc.empty else ""
    if d.empty: return [], []
    Fv, df1, df2, p, et = cell(d, 0, "F"), cell(d, 0, "df"), cell(d, 1, "df"), cell(d, 0, "p"), cell(d, 0, "partial_eta_sq")
    post = ""
    if not ph.empty:
        sigp = ph[ph["p_adjusted"] < .05]
        if len(sigp): post = "事后比较（Bonferroni）显著的配对：" + "；".join(f"{r['pair1']} vs {r['pair2']}" for _, r in sigp.iterrows()) + "。"
    zh = [f"单因素方差分析检验自变量 {iv} 的各水平在因变量 {dv} 上的均值是否相等。{groups}",
          f"组间效应：F({int(df1)}, {int(df2)}) = {num(Fv)}, {result_report(p)}, η²p = {num(et, 3)}。",
          concl(p, sig_txt=f"结论：至少有两组均值不同（{ES_ETA(et)}效应），自变量解释了因变量方差的 {num(et*100, 1)}%。{post}",
                ns_txt="结论：各组均值无显著差异，不应再报告事后比较。")]
    en = [f"One-way ANOVA tests whether levels of {iv} differ on {dv}. {groups}",
          f"Between-group effect: F({int(df1)}, {int(df2)}) = {num(Fv)}, {result_report(p)}, partial η² = {num(et, 3)}.",
          concl(p, sig_txt=f"Conclusion: at least two groups differ; the factor explains {num(et*100, 1)}% of variance. {post}",
                ns_txt="Conclusion: no significant group differences; post hoc tests should not be reported.")]
    # 方差齐性：Levene 显著时不得只报合并方差的 F —— 必须读稳健检验表并给出明确警示。
    lev = g1(f, "one_way_anova_levene"); rob = g1(f, "one_way_anova_robust_tests")
    lev_p = None
    if not lev.empty:
        for col in ("p", "levene_p"):
            if col in lev.columns:
                lev_p = cell(lev, 0, col); break
    if isinstance(lev_p, (int, float)) and not pd.isna(lev_p) and lev_p < .05:
        robust_txt_zh = robust_txt_en = ""
        if not rob.empty:
            try:
                rrow = rob.iloc[0]
                parts = [f"{c} = {num(rrow[c])}" for c in rob.columns
                         if c.lower() not in ("test", "note", "method") and pd.notna(rrow[c])]
                if parts:
                    robust_txt_zh = "稳健检验（Welch / Brown-Forsythe）：" + "，".join(parts[:5]) + "。"
                    robust_txt_en = "Robust tests (Welch / Brown-Forsythe): " + ", ".join(parts[:5]) + ". "
            except Exception:
                pass
        zh.append(f"注意：Levene 方差齐性检验 p = {num(lev_p)} < .05，**方差不齐**，上表合并方差的 F 检验前提不成立；"
                  f"应改报 Welch 校正结果或改用非参数检验（Kruskal-Wallis）。{robust_txt_zh}")
        en.append(f"Note: Levene's test p = {num(lev_p)} < .05 (unequal variances); the pooled F above is not valid. "
                  f"Report the Welch-corrected test or use a nonparametric method (Kruskal-Wallis). {robust_txt_en}")
    return zh, en

def s_two_way_anova(f):
    d = g1(f, "two_way_anova_anova"); om = g1(f, "two_way_anova_simple_effects_omnibus")
    if d.empty: return [], []
    Fv, df1, df2, p, et = cell(d, 2, "F"), cell(d, 2, "df"), cell(d, 3, "df"), cell(d, 2, "p"), cell(d, 2, "partial_eta_sq")
    omni = ""
    if not om.empty:
        sig = om[om["p"] < .05]
        if len(sig): omni = "显著的简单效应（总检验）：" + "；".join(f"{r['方向']}（{result_report(r['p'])}）" for _, r in sig.iterrows()) + "。"
    zh = [f"两因素方差分析同时检验两个主效应与一个交互效应。",
          f"交互效应：F({int(df1)}, {int(df2)}) = {num(Fv)}, {result_report(p)}, η²p = {num(et, 3)}。",
          concl(p, sig_txt=f"结论：交互显著（{ES_ETA(et)}效应），主效应失去直接解释意义，须以简单效应（表：简单效应总检验）解释；{omni}",
                ns_txt="结论：交互不显著，分别解释两个主效应即可（主效应显著且>2水平时做事后比较）。")]
    en = [f"Two-way ANOVA tests two main effects and one interaction.",
          f"Interaction: F({int(df1)}, {int(df2)}) = {num(Fv)}, {result_report(p)}, partial η² = {num(et, 3)}.",
          concl(p, sig_txt="Conclusion: the interaction is significant; interpret effects through simple effects.",
                ns_txt="Conclusion: the interaction is not significant; interpret the main effects separately.")]
    return zh, en

def s_rm_anova(f):
    m = g1(f, "rm_anova_mauchly"); d = g1(f, "rm_anova_within_tests")
    if d.empty: return [], []
    pm = cell(m, 0, "p")
    row = 0 if (pm is None or pm >= .05) else 1
    Fv, df1, df2, p, et = cell(d, row, "F"), cell(d, row, "df1"), cell(d, row, "df2"), cell(d, row, "p"), cell(d, row, "partial_eta_sq")
    zh = [f"重复测量方差分析检验同一批被试在多个时间点/条件下的均值变化。",
          f"Mauchly检验 {result_report(pm)}；时间效应：F({num(df1)}, {num(df2)}) = {num(Fv)}, {result_report(p)}, η²p = {num(et, 3)}。",
          concl(p, sig_txt="结论：随时间/条件变化显著；配对事后比较（Bonferroni）定位了具体时间点差异。",
                ns_txt="结论：各时间点均值无显著差异。")]
    en = [f"Repeated-measures ANOVA tests mean changes across time/conditions within subjects.",
          f"Mauchly's test {result_report(pm)}; time effect: F({num(df1)}, {num(df2)}) = {num(Fv)}, {result_report(p)}, partial η² = {num(et, 3)}.",
          concl(p, sig_txt="Conclusion: significant change over time; see Bonferroni pairwise comparisons.",
                ns_txt="Conclusion: no significant change across time points.")]
    return zh, en

def s_mixed_anova(f):
    d = g1(f, "mixed_anova_anova")
    if d.empty: return [], []
    Fv, p, et = cell(d, 2, "F"), cell(d, 2, "p"), cell(d, 2, "partial_eta_sq")
    zh = [f"混合设计方差分析含一个被试间因子与一个被试内因子。",
          f"组别×时间交互：F = {num(Fv)}, {result_report(p)}, η²p = {num(et, 3)}。",
          concl(p, sig_txt="结论：交互显著，处理效果取决于\"组别×时间\"组合，见简单效应表（多变量Pillai行对应SPSS EM均值输出）。",
                ns_txt="结论：交互不显著，分别解释两个主效应。")]
    en = [f"Mixed-design ANOVA includes one between- and one within-subject factor.",
          f"Group × time interaction: F = {num(Fv)}, {result_report(p)}, partial η² = {num(et, 3)}.",
          concl(p, sig_txt="Conclusion: significant interaction; see the simple-effects tables.",
                ns_txt="Conclusion: no significant interaction; interpret the main effects.")]
    return zh, en

def s_ancova(f):
    d = g1(f, "ancova_ancova")
    if d.empty: return [], []
    Fv, df1, df2, p, et = cell(d, 1, "F"), cell(d, 1, "df"), cell(d, 2, "df"), cell(d, 1, "p"), cell(d, 1, "partial_eta_sq")
    zh = [f"协方差分析在控制连续协变量后比较各组（报告调整后均值而非原始均值）。",
          f"控制协变量后的组间效应：F({int(df1)}, {int(df2)}) = {num(Fv)}, {result_report(p)}, η²p = {num(et, 3)}。",
          concl(p, sig_txt="结论：控制协变量后组间差异仍显著；比较各组调整后均值与成对比较结果。",
                ns_txt="结论：控制协变量后组间无显著差异。")]
    en = [f"ANCOVA compares groups after adjusting for a continuous covariate (adjusted means are reported).",
          f"Adjusted group effect: F({int(df1)}, {int(df2)}) = {num(Fv)}, {result_report(p)}, partial η² = {num(et, 3)}.",
          concl(p, sig_txt="Conclusion: group differences remain after adjustment; compare adjusted means.",
                ns_txt="Conclusion: no group differences after adjustment.")]
    return zh, en

def s_correlation(f):
    d = g1(f, "correlation_correlations"); pc = g1(f, "correlation_partial_correlation")
    if d.empty: return [], []
    sig = d[(d["method"] == "pearson") & (d["p"] < .05)]
    zh = [f"相关分析考察连续变量间的关联（Pearson为主，偏态/等级用Spearman）。",
          ("显著相关：" + "；".join(f"{r['var1']}–{r['var2']}：r = {num(r['r'])}（{result_report(r['p'])}，{ES_R(r['r'])}）" for _, r in sig.iterrows()) + "。") if len(sig) else "Pearson 相关均未达显著。",
          "结论：相关显著只说明关联存在，不等于因果；r² 表示共变方差比例。偏相关表给出控制第三变量后的独特关联。"]
    en = [f"Correlation analysis examines associations among continuous variables.",
          ("Significant: " + "; ".join(f"{r['var1']}–{r['var2']}: r = {num(r['r'])} ({result_report(r['p'])})" for _, r in sig.iterrows()) + ".") if len(sig) else "No significant Pearson correlations.",
          "Conclusion: correlation is not causation; the partial-correlation table shows unique association controlling for a third variable."]
    return zh, en

def s_regression(f):
    m = g1(f, "regression_model_summary"); d = g1(f, "regression_anova")
    if m.empty: return [], []
    r2, ar2, Fv, p = cell(m, 0, "R_squared"), cell(m, 0, "adjusted_R_squared"), cell(d, 0, "F"), cell(d, 0, "p")
    zh = [f"多元线性回归用多个自变量预测连续因变量（Enter法整体进入）。",
          f"整体模型检验：F = {num(Fv)}, {result_report(p)}；R² = {num(r2, 3)}（调整R² = {num(ar2, 3)}），模型解释因变量方差的 {num(r2*100, 1)}%。",
          "结论：各预测变量的独特贡献看系数表（B、β、t、p、VIF）；层次增量表给出每个变量额外解释的 ΔR²。"]
    en = [f"Multiple linear regression predicts a continuous outcome from several predictors (Enter method).",
          f"Overall model: F = {num(Fv)}, {result_report(p)}; R² = {num(r2, 3)} (adj. R² = {num(ar2, 3)}), explaining {num(r2*100, 1)}% of the variance.",
          "Conclusion: unique contributions are in the coefficients table; the hierarchical table gives ΔR² per added predictor."]
    return zh, en

def s_moderation(f):
    d = g1(f, "moderation_interaction_model"); sl = g1(f, "moderation_simple_slopes")
    if d.empty: return [], []
    b3 = p3 = None
    for _, r in d.iterrows():
        if str(r["term"]) == "Xc:Zc": b3, p3 = float(r["B"]), float(r["p"])
    s1 = s3 = None
    p1 = p3s = None
    # 这 4 个变量必须一起初始化：简单斜率表缺失/为空时不进 `if not sl.empty` 分支，
    # 而下面的句子**无条件**引用 p1/p3s → UnboundLocalError，会让整份报告生成失败
    # （内测真实发生：跑"调节效应"时 moderation_simple_slopes.csv 缺失，HTML 报告整份出不来）
    if not sl.empty:
        s1, s3 = cell(sl, 0, "simple_slope_of_X"), cell(sl, 2, "simple_slope_of_X")
        p1, p3s = cell(sl, 0, "p"), cell(sl, 2, "p")
    slopes_zh = (f"简单斜率：Z低(−1SD)时 = {num(s1, 3)}（{result_report(p1)}），Z高(+1SD)时 = {num(s3, 3)}（{result_report(p3s)}）。"
                 if not sl.empty else
                 "简单斜率表（moderation_simple_slopes.csv）未生成或为空，本节无法给出各条件下的效应；"
                 "交互项系数与上方结论不受影响，请检查该方法是否运行完整。")
    slopes_en = (f"simple slopes: low Z = {num(s1, 3)} ({result_report(p1)}), high Z = {num(s3, 3)} ({result_report(p3s)})."
                 if not sl.empty else
                 "The simple-slopes table (moderation_simple_slopes.csv) is missing or empty, so conditional "
                 "effects cannot be reported; the interaction term above is unaffected.")
    zh = [f"调节分析检验\"X的效应是否随Z变化\"（中心化乘积项，PROCESS Model 1）。",
          f"交互项 B = {num(b3, 3)}, {result_report(p3)}；{slopes_zh}",
          concl(p3, sig_txt="结论：调节效应显著，X的效应强度随Z变化；按简单斜率与交互图描述各条件下的效应。",
                ns_txt="结论：无调节证据，X的效应不随Z变化。")]
    en = [f"Moderation analysis tests whether the effect of X varies with Z (centered product term, PROCESS Model 1).",
          f"Interaction B = {num(b3, 3)}, {result_report(p3)}; {slopes_en}",
          concl(p3, sig_txt="Conclusion: significant moderation; describe the effect at each level of Z.",
                ns_txt="Conclusion: no evidence of moderation.")]
    return zh, en

def s_chi_gof(f):
    d = g1(f, "chi_square_gof_test")
    if d.empty: return [], []
    c, df_, p, N = cell(d, 0, "chisq"), cell(d, 0, "df"), cell(d, 0, "p"), cell(d, 0, "N")
    zh = [f"卡方适合度检验单个分类变量的分布是否符合理论比例。",
          f"χ²({int(df_)}, N = {int(N)}) = {num(c)}, {result_report(p)}。",
          concl(p, sig_txt="结论：观察分布显著偏离理论比例；残差最大的类别是主要来源。",
                ns_txt="结论：观察分布与理论比例无显著差异。")]
    en = [f"The chi-square goodness-of-fit test compares observed frequencies with theoretical proportions.",
          f"χ²({int(df_)}, N = {int(N)}) = {num(c)}, {result_report(p)}.",
          concl(p, sig_txt="Conclusion: the distribution deviates from the theoretical proportions.",
                ns_txt="Conclusion: no significant deviation.")]
    return zh, en

def s_chi_ind(f):
    d = g1(f, "chi_square_independence_tests"); e = g1(f, "chi_square_independence_effect_size")
    if d.empty: return [], []
    c, df_, p, N = cell(d, 0, "chisq"), cell(d, 0, "df"), cell(d, 0, "p"), cell(d, 0, "N")
    eff = cell(e, 0, "value")
    zh = [f"卡方独立性检验两个分类变量是否关联。",
          f"χ²({int(df_)}, N = {int(N)}) = {num(c)}, {result_report(p)}, Cramér's V = {num(eff, 3)}。",
          concl(p, sig_txt="结论：两变量关联显著（效应量见上）；用调整标准化残差定位关键格子。",
                ns_txt="结论：两变量相互独立（无显著关联）。")]
    en = [f"The chi-square test of independence assesses association between two categorical variables.",
          f"χ²({int(df_)}, N = {int(N)}) = {num(c)}, {result_report(p)}, Cramér's V = {num(eff, 3)}.",
          concl(p, sig_txt="Conclusion: significant association; locate driving cells with adjusted standardized residuals.",
                ns_txt="Conclusion: the two variables are independent.")]
    return zh, en

def s_mann_whitney(f):
    d = g1(f, "mann_whitney_test"); rk = g1(f, "mann_whitney_ranks")
    if d.empty: return [], []
    U, Z, p, r = cell(d, 0, "Mann_Whitney_U"), cell(d, 0, "Z"), cell(d, 0, "p_asymptotic_2tailed"), cell(d, 0, "effect_size_r")
    mds = ""
    if not rk.empty: mds = "；".join(f"{row['group']}: Mdn = {num(row['median'], 1)}" for _, row in rk.iterrows())
    zh = [f"Mann-Whitney U 检验（两组独立、偏态/等级数据；报告中位数）。{mds}",
          f"U = {num(U)}, Z = {num(Z)}, {result_report(p)}, r = {num(r)}。",
          concl(p, sig_txt="结论：两组分布位置显著不同（比较中位数方向）。", ns_txt="结论：两组无显著差异。")]
    en = [f"Mann-Whitney U test (two independent groups, skewed/ordinal data; report medians). {mds.replace('；', '; ')}",
          f"U = {num(U)}, Z = {num(Z)}, {result_report(p)}, r = {num(r)}.",
          concl(p, sig_txt="Conclusion: the two groups differ in location.", ns_txt="Conclusion: no significant difference.")]
    return zh, en

def s_wilcoxon(f):
    d = g1(f, "wilcoxon_signed_test")
    if d.empty: return [], []
    dlab = str(d.iloc[0].get("difference_direction", "后−前"))
    W, Z, p, r = cell(d, 0, "Wilcoxon_W"), cell(d, 0, "Z"), cell(d, 0, "p_asymptotic_2tailed"), cell(d, 0, "effect_size_r")
    zh = [f"Wilcoxon 符号秩检验（前后测、差值偏态/等级；差值 = {dlab}）。",
          f"W = {num(W)}, Z = {num(Z)}, {result_report(p)}, r = {num(r)}。",
          concl(p, sig_txt="结论：前后测差异显著；报告中位数与方向。", ns_txt="结论：前后测无显著差异。")]
    en = [f"Wilcoxon signed-rank test (pre-post, skewed/ordinal differences; difference = {dlab}).",
          f"W = {num(W)}, Z = {num(Z)}, {result_report(p)}, r = {num(r)}.",
          concl(p, sig_txt="Conclusion: significant pre-post difference.", ns_txt="Conclusion: no significant difference.")]
    return zh, en

def s_kruskal(f):
    d = g1(f, "kruskal_wallis_test")
    if d.empty: return [], []
    H, df_, p, e2 = cell(d, 0, "Kruskal_Wallis_H"), cell(d, 0, "df"), cell(d, 0, "p"), cell(d, 0, "epsilon_squared")
    zh = [f"Kruskal-Wallis H 检验（≥3组独立、偏态/等级数据）。",
          f"H({int(df_)}) = {num(H)}, {result_report(p)}, ε² = {num(e2, 3)}。",
          concl(p, sig_txt="结论：至少两组分布位置不同；Bonferroni事后表定位差异组。", ns_txt="结论：各组无显著差异。")]
    en = [f"Kruskal-Wallis H test (≥3 independent groups, skewed/ordinal data).",
          f"H({int(df_)}) = {num(H)}, {result_report(p)}, ε² = {num(e2, 3)}.",
          concl(p, sig_txt="Conclusion: at least two groups differ; see Bonferroni pairwise table.", ns_txt="Conclusion: no significant differences.")]
    return zh, en

def s_friedman(f):
    d = g1(f, "friedman_test")
    if d.empty: return [], []
    C, df_, p, W = cell(d, 0, "Friedman_chisq"), cell(d, 0, "df"), cell(d, 0, "p"), cell(d, 0, "kendalls_W")
    zh = [f"Friedman 检验（≥3次重复测量、偏态/等级数据）。",
          f"χ²({int(df_)}) = {num(C)}, {result_report(p)}, Kendall's W = {num(W)}。",
          concl(p, sig_txt="结论：各条件评分差异显著；事后成对比较定位差异。", ns_txt="结论：各条件无显著差异。")]
    en = [f"Friedman test (≥3 related conditions, skewed/ordinal data).",
          f"χ²({int(df_)}) = {num(C)}, {result_report(p)}, Kendall's W = {num(W)}.",
          concl(p, sig_txt="Conclusion: conditions differ significantly; see pairwise table.", ns_txt="Conclusion: no significant differences.")]
    return zh, en

def s_mediation(f):
    d = g1(f, "mediation_indirect")
    if d.empty: return [], []
    ab, lo, hi = cell(d, 0, "value"), cell(d, 2, "value"), cell(d, 3, "value")
    ratio = cell(d, 4, "value")
    # CI 缺失（NaN/None）与"含 0"是两回事：前者是"算不出"，绝不能说成"无中介效应"。
    ci_available = (isinstance(lo, (int, float)) and isinstance(hi, (int, float))
                    and not pd.isna(lo) and not pd.isna(hi))
    sig = ci_available and (lo > 0 or hi < 0)
    if not ci_available:
        ci_txt_zh = "Boot 95%CI 无法计算（Bootstrap 重抽样出现退化，通常是样本量过小或变量近乎常量）→ 不能据此判断中介是否显著，请检查数据后重跑。"
        ci_txt_en = "The boot 95% CI could not be computed (degenerate bootstrap resamples - usually too few cases or a near-constant variable); mediation cannot be judged from this output."
        zh = ["中介分析检验\"X是否通过M影响Y\"（三变量模型，Bootstrap 5000 次百分位CI）。",
              f"间接效应 ab = {num(ab, 3)}；{ci_txt_zh}", ci_txt_zh]
        en = ["Mediation analysis tests whether X affects Y through M (bootstrap 5,000, percentile CI).",
              f"Indirect effect ab = {num(ab, 3)}. {ci_txt_en}", ci_txt_en]
        return zh, en
    zh = [f"中介分析检验\"X是否通过M影响Y\"（三变量模型，Bootstrap 5000 次百分位CI）。",
          f"间接效应 ab = {num(ab, 3)}，Boot 95%CI [{num(lo, 3)}, {num(hi, 3)}]{'（不含0，中介显著）' if sig else '（含0，中介不显著）'}，占总效应 {num(100*ratio, 1) if ratio else 'NA'}%。",
          "结论：" + ("间接效应显著；直接效应c'显著为部分中介、不显著为（接近）完全中介，建议同时报告ab、CI与占比。" if sig else "未发现显著的中介效应（CI 含 0）。")]
    en = [f"Mediation analysis tests whether X affects Y through M (bootstrap 5,000, percentile CI).",
          f"Indirect effect ab = {num(ab, 3)}, boot 95% CI [{num(lo, 3)}, {num(hi, 3)}]{' (significant)' if sig else ' (not significant)'}, {num(100*ratio, 1) if ratio else 'NA'}% of total effect.",
          "Conclusion: " + ("The indirect effect is significant; interpret partial vs. full mediation via the direct effect c'." if sig else "No significant mediation (CI includes 0).")]
    return zh, en

def s_power(f):
    d = g1(f, "power_sample_size_table")
    if d.empty: return [], []
    r = d[d["effect"].astype(str).str.contains("d = 0.5") & (d["power"] == 0.8)]
    n = int(r.iloc[0]["N_total"] / 2) if len(r) else None
    zh = [f"功效分析用于开题/设计阶段估算样本量（对应G*Power）。",
          f"参考：中等效应 d = 0.50、α = .05（双尾）、power = .80 需每组 {n} 人。",
          "结论：小效应需要的样本量远大于中等/大效应；\"不显著\"可能是功效不足而非无效应。"]
    en = [f"Power analysis estimates required sample sizes (G*Power equivalent).",
          f"Reference: detecting d = 0.50 with power = .80 requires {n} per group.",
          "Conclusion: small effects need much larger samples; a non-significant result may reflect low power."]
    return zh, en

# ── module registry ──

def s_datacheck(folder: Path):
    """数据准备（第 0 步）：表格由管线写在上方 prep 段，这里只出叙事文字。"""
    def g(name):
        pth = folder / f"00_prep_{name}.csv"
        return pd.read_csv(pth, encoding="utf-8-sig") if pth.exists() else pd.DataFrame()
    ma = g("missing_audit"); imp = g("imputation"); os_ = g("outlier_summary")
    outl = g("outliers"); integ = g("integrity"); summ = g("preparation_summary")

    zh = []
    if not ma.empty:
        miss = ma[ma["N_missing"] > 0]
        if miss.empty:
            zh.append("缺失值审计：各变量均无缺失值。")
        else:
            zh.append("缺失值审计：" + "；".join(
                f"{r['variable']} 缺 {int(r['N_missing'])} 个（{r['pct_missing']}%）" for _, r in miss.iterrows()) + "。")
    if not imp.empty:
        tot = int(imp["n_imputed"].sum())
        zh.append(f"缺失值处理：对 {len(imp)} 个变量做了均值插补，共替换 {tot} 个缺失值。"
                  "注意均值插补会缩小该变量的方差并削弱它与其他变量的相关，报告时应写明处理方式与个数。")
    if not os_.empty:
        nz = int(os_["n_z"].sum()); ni = int(os_["n_iqr"].sum()); ne = int(os_["n_extreme_iqr"].sum())
        zh.append(f"异常值筛查：按 |Z| 规则命中 {nz} 个个案次，按 1.5×IQR 命中 {ni} 个个案次（含超过 3×IQR 的极端值 {ne} 个）。"
                  "两条规则看的是同一批可疑个案但判定标准不同：Z 以标准差为单位（小样本有理论上限 (n−1)/√n，会漏报），"
                  "1.5×IQR 以四分位距为单位、不假设正态。异常值不一定是错误，删除前请回看原始作答记录。")
    if not outl.empty:
        zh.append("异常个案明细：" + "；".join(
            f"{r['variable']} 的 {r['id']}（值 {r['value']}，{r['reason']}）" for _, r in outl.head(8).iterrows()) +
            ("…" if len(outl) > 8 else "") + "。")
    if not integ.empty and not (integ["check"] == "未发现问题").all():
        zh.append("数据完整性检查：" + "；".join(
            f"{r['check']}——{r['variable']}：{r['detail']}" for _, r in integ.iterrows()) + "。")
    if not summ.empty:
        vals = {r["item"]: r["value"] for _, r in summ.iterrows()}
        zh.append("数据准备摘要（写论文方法部分可直接用这段）：原始 N = "
                  f"{int(vals.get('原始个案数', 0))}；均值插补 {int(vals.get('均值插补的缺失值个数', 0))} 个缺失值；"
                  f"按 |Z| 规则删除 {int(vals.get('按 |Z| 规则删除的个案数', 0))} 个个案；"
                  f"最终用于分析的 N = {int(vals.get('最终个案数', 0))}。")
    if not zh:
        zh = ["数据准备：未产生结果表（可能本次输入为模拟数据或数据准备被跳过）。"]
    en = [("Data preparation: see the tables in section 0 (missing-value audit, outlier screening by |Z| and 1.5*IQR, integrity checks). "
           "Mean imputation reduces the variance of the imputed variable; report the method and the number of imputed values. "
           "Outliers are not necessarily errors - inspect the raw responses before removing them.")]
    return zh, en

def _guarded(fn, mid):
    """把一个方法的中文/英文报告段落生成器包成"永不抛异常"的版本。

    动机（内测真实故障）：某个方法的结果表缺失或为空时，段落生成器可能抛异常，
    而它是在**整份报告**的循环里被调用的——于是一处小缺失会让 Word 与 HTML 报告**整份**出不来，
    用户看到的是一个与真正原因相距很远的 traceback。
    现在改为：只跳过该节的文字，插入一条显式说明（报告里看得见），并把原始 traceback 打到 stderr，
    统计表本身照常输出。

    注意：这里只捕获 Exception（不捕获 SystemExit/KeyboardInterrupt），也不吞掉 KeyboardInterrupt。
    """
    def wrapped(folder):
        try:
            return fn(folder)
        except Exception as e:                                   # noqa: BLE001 —— 故意兜住一切报告层异常
            import traceback as _tb
            print(f"[警告] 方法 {mid} 的报告段落生成失败，已跳过该节文字（统计表不受影响）：{e}", file=sys.stderr)
            print(_tb.format_exc(), file=sys.stderr)
            zh = [f"（本节报告段落生成失败：{type(e).__name__}: {e}）",
                  "该方法的统计表仍已正常输出，可直接查阅上方表格。"
                  "请把本结果目录连同 run_log.txt 反馈给作者以便定位。"]
            en = [f"(Report paragraphs for this section failed to generate: {type(e).__name__}: {e}.)",
                  "The statistical tables for this method are still available above. "
                  "Please report this folder together with run_log.txt."]
            return zh, en
    wrapped.__name__ = getattr(fn, "__name__", str(mid))
    return wrapped

def build_registry():
    handlers = {
        "datacheck": s_datacheck,
        "descriptives": s_descriptives, "normality": s_normality, "one_sample_t": s_one_sample_t,
        "independent_t": s_independent_t, "paired_t": s_paired_t, "one_way_anova": s_one_way_anova,
        "two_way_anova": s_two_way_anova, "rm_anova": s_rm_anova, "mixed_anova": s_mixed_anova,
        "ancova": s_ancova, "correlation": s_correlation, "regression": s_regression,
        "moderation": s_moderation, "chi_square_gof": s_chi_gof, "chi_square_independence": s_chi_ind,
        "mann_whitney": s_mann_whitney, "wilcoxon_signed": s_wilcoxon, "kruskal_wallis": s_kruskal,
        "friedman": s_friedman, "mediation": s_mediation, "power": s_power,
    }
    # 统一包上保护层：单个方法的段落生成失败不再拖垮整份报告（Word 与 HTML 两条路径共用本注册表）
    return {mid: _guarded(fn, mid) for mid, fn in handlers.items()}

# table captions for known outputs (pattern suffix -> zh/en captions); unknown files fall back to file name
CAPTIONS = {
    "desc": ("前置描述统计", "Preliminary descriptives"), "norm": ("前置正态性检验", "Preliminary normality tests"),
    "stats": ("描述统计/检验统计", "Descriptive/test statistics"), "normality_tests": ("正态性检验", "Normality tests"),
    "levene": ("Levene方差齐性检验", "Levene's test"), "group_stats": ("组统计量", "Group statistics"),
    "test": ("检验统计", "Test statistics"), "correlation": ("配对相关/相关", "Correlation"),
    "descriptives": ("描述统计", "Descriptives"), "anova": ("方差分析表", "ANOVA table"),
    "posthoc_lsd": ("事后比较（LSD）", "Post hoc (LSD)"), "posthoc_tukey": ("事后比较（Tukey）", "Post hoc (Tukey)"),
    "posthoc_bonferroni": ("事后比较（Bonferroni）", "Post hoc (Bonferroni)"), "robust_tests": ("稳健检验", "Robust tests"),
    "cell_descriptives": ("单元格描述统计", "Cell descriptives"), "simple_effects_omnibus": ("简单效应总检验（双向）", "Simple effects (omnibus, both directions)"),
    "simple_effects_a_by_b": ("成对比较：A在各B水平", "Pairwise: A within B"), "simple_effects_b_by_a": ("成对比较：B在各A水平", "Pairwise: B within A"),
    "emmeans": ("估算边际均值", "Estimated marginal means"), "mauchly": ("Mauchly球形检验", "Mauchly's test"),
    "within_tests": ("被试内效应检验", "Within-subjects effects"), "posthoc_paired": ("配对比较（Bonferroni）", "Pairwise comparisons (Bonferroni)"),
    "simple_time_by_group": ("简单效应：各组内时间效应", "Simple effects of time within groups"),
    "simple_group_by_time": ("简单效应：各时间点组间差异", "Simple effects of group at times"),
    "slope_homogeneity": ("斜率同质性检验", "Homogeneity of regression slopes"),
    "levene_errorvar": ("误差方差齐性（Levene，模型残差）", "Levene's test (model residuals)"),
    "adjusted_means": ("调整后均值", "Adjusted means"), "posthoc_adjusted": ("调整后组间比较（Bonferroni）", "Adjusted pairwise (Bonferroni)"),
    "correlations": ("两两相关", "Pairwise correlations"), "partial_correlation": ("偏相关", "Partial correlation"),
    "model_summary": ("模型汇总", "Model summary"), "hierarchical": ("层次进入增量检验", "Hierarchical increment tests"),
    "coefficients": ("回归系数", "Coefficients"), "residuals_stats": ("残差统计", "Residuals statistics"),
    "interaction_model": ("调节（交互）模型", "Moderation (interaction) model"), "simple_slopes": ("简单斜率", "Simple slopes"),
    "frequencies": ("频数表", "Frequencies"), "tests": ("检验汇总", "Tests summary"), "effect_size": ("效应量", "Effect size"),
    "crosstab": ("交叉表（观察/期望/残差）", "Crosstab (observed/expected/residuals)"),
    "ranks": ("秩统计", "Ranks"), "posthoc_pairwise": ("事后两两比较（Bonferroni）", "Pairwise (Bonferroni)"),
    "paths": ("中介路径系数", "Mediation paths"), "indirect": ("间接效应与Bootstrap检验", "Indirect effect and bootstrap test"),
    "sample_size_table": ("所需样本量", "Required sample sizes"), "zscores": ("Z分数", "Z scores"),
}

def failed_methods_note(folder: Path, zh: bool):
    """读 run_log.txt 的 Failed methods 行（stats_pipeline 单方法失败但整次运行继续时写入），
    让 Word 报告也标注哪些方法未能完成；无失败时返回 None。"""
    log = folder / "run_log.txt"
    if not log.exists():
        return None
    for line in log.read_text(encoding="utf-8", errors="replace").splitlines():
        if line.startswith("Failed methods:"):
            names = line.split(":", 1)[1].strip()
            if names and names.lower() != "none":
                return (f"注意：{names} 未能完成（原因与建议见运行日志及 stats_report_zh.md 中的对应标注）；本报告仅包含成功完成的方法。"
                        if zh else
                        f"Note: the following methods did not complete: {names}. See the run log and stats_report_zh.md for reasons; this report covers the completed methods only.")
    return None

def present_methods(folder: Path):
    reg = build_registry()
    out = []
    # 第 0 步数据准备的产物用 00_prep_* 前缀（不占方法编号），单独认成第 0 节排在最前
    if glob.glob(str(folder / "00_prep_*.csv")):
        out.append((0, "datacheck"))
    for mid in reg:
        if mid == "datacheck":
            continue
        hits = sorted(glob.glob(str(folder / f"??_{mid}_*.csv")))
        hits = [h for h in hits if not h.endswith("_data.csv")]
        if hits:
            m = re.match(r"(\d\d)_", os.path.basename(hits[0]))
            out.append((int(m.group(1)) if m else 99, mid))
    return sorted(out)

def method_tables(folder: Path, mid: str):
    pat = "00_prep_*.csv" if mid == "datacheck" else f"??_{mid}_*.csv"
    files = [Path(h) for h in sorted(glob.glob(str(folder / pat))) if not h.endswith("_data.csv")]
    out = []
    used = set()
    for pth in files:
        stem = pth.name.split("_", 2)[2].replace(".csv", "")
        cap = None
        for key, (z, e) in CAPTIONS.items():
            if stem == key or stem.endswith("_" + key):
                cap = (z, e); break
        if cap is None:
            for key, (z, e) in CAPTIONS.items():
                if key in stem: cap = (z, e); break
        if cap is None: cap = (stem, stem)
        if cap in used: cap = (cap[0] + "（续）", cap[1] + " (cont.)")
        used.add(cap)
        out.append((pth, cap))
    return out

FIG_SUFFIXES = {"interaction": ("交互图", "Interaction plot"), "profile": ("轮廓图", "Profile plot"),
                "diagnostics": ("残差诊断图", "Residual diagnostics"), "qqplot": ("Q-Q图", "Q-Q plot"),
                "histogram": ("直方图", "Histogram"), "boxplot": ("箱线图", "Boxplot"),
                "scatter": ("散点图", "Scatterplot"), "means": ("均值图", "Means plot"),
                "ancova": ("ANCOVA回归线", "ANCOVA regression lines"), "path_diagram": ("中介路径图", "Mediation path diagram")}

def build_report(folder: Path, lang: str):
    reg = build_registry()
    doc = Document()
    normal = doc.styles["Normal"]; normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei"); normal.font.size = Pt(10.5)
    zh = lang == "zh"
    title = doc.add_heading("Psychostat 心理统计分析报告" if zh else "Psychostat Statistical Analysis Report ", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("自动生成初稿｜每节含文字说明（方法—数据—结果—结论）与全部结果表（三线表）；请核对后用于正式报告。统计实现与IBM SPSS默认参数一致。"
                      if zh else
                      "Auto-generated draft; each section contains narrative (method - data - results - conclusion) and all result tables in three-line style. Computations follow IBM SPSS defaults.",
                      style="Subtitle")
    methods = present_methods(folder)
    doc.add_heading("摘要" if zh else "Summary", 1)
    names = "、".join(f"{m[1]}" for m in methods) if zh else ", ".join(m[1] for m in methods)
    doc.add_paragraph((f"本报告基于 {len(methods)} 个统计模块自动生成：{names}。各模块的数据文件（NN_方法_data.csv，UTF-8-BOM）可直接导入SPSS复现；"
                       "每节先给文字说明（这个方法做什么、数据长什么样、结果如何、结论是什么），随后附上该方法产生的全部表格与图形。"
                       if zh else
                       f"This report covers {len(methods)} modules: {names}. Each section first narrates the method, data, results and conclusion, then embeds every table and figure produced by the module."))
    fail_note = failed_methods_note(folder, zh)
    if fail_note:
        doc.add_paragraph(fail_note)
    tab_no = 0
    for idx, mid in methods:
        doc.add_heading(f"{idx}. {mid}", 1)
        zh_paras, en_paras = reg[mid](folder)
        paras = zh_paras if zh else en_paras
        for p in paras:
            if p: doc.add_paragraph(p)
        for pth, (cz, ce) in method_tables(folder, mid):
            tab_no += 1
            three_line_table(doc, read_csv(folder, pth.name),
                             (f"表{tab_no} {cz}（{mid}）" if zh else f"Table {tab_no} {ce} ({mid})"),
                             p_matrix="matrix_p_" in pth.stem)
        for sfx, (fz, fe) in FIG_SUFFIXES.items():
            fig_pat = f"00_prep_{sfx}.png" if mid == "datacheck" else f"??_{mid}_{sfx}.png"
            for hit in sorted(glob.glob(str(folder / fig_pat))):
                tab_no += 1
                figure(doc, Path(hit), f"图{tab_no} {fz}（{mid}）" if zh else f"Figure {tab_no} {fe} ({mid})")
    return doc

def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--result-dir-utf8-base64", required=True); args = ap.parse_args()
    folder = Path(base64.b64decode(args.result_dir_utf8_base64).decode("utf-8"))
    build_report(folder, "zh").save(str(folder / "stats_report_zh.docx"))
    build_report(folder, "en").save(str(folder / "stats_report_en.docx"))
    print("REPORTS_OK")

if __name__ == "__main__":
    main()
