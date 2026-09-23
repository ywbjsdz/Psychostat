"""把 docs/功能速览.md 渲染为项目根目录的《Psychostat 功能速览.docx》。

为什么单独做一份：完整说明书 878 行，同学多半不看；速览目标 3 分钟读完，
只回答"这是什么 / 我该用哪条线 / 这些术语啥意思 / 最容易踩什么坑"。
复用 merge_manuals 的 markdown 渲染器，排版与说明书一致。
用法（Windows）：py -3 scripts/make_quickstart_docx.py
"""
from __future__ import annotations
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
from docx import Document                                   # noqa: E402
from merge_manuals import add_md, style_doc, save_docx       # noqa: E402


def main() -> int:
    src = ROOT / "docs" / "功能速览.md"
    if not src.exists():
        print(f"找不到源文件：{src}")
        return 1
    doc = Document()
    style_doc(doc)
    add_md(doc, src.read_text(encoding="utf-8"), base_level=1)
    out = ROOT / "Psychostat 功能速览.docx"
    save_docx(doc, out)
    print(f"已生成：{out}（{out.stat().st_size} 字节，段落 {len(doc.paragraphs)}，表格 {len(doc.tables)}）")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
