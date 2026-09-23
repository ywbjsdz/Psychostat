"""Merge the three Psychostat manuals (stats/CTT/IRT) into one Word document."""
from __future__ import annotations
import sys
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parent.parent


def save_docx(doc, path) -> None:
    """先写到同目录临时文件再原子替换：避免在部分文件系统上直接原地保存失败。"""
    import os
    path = str(path)
    tmp = path + ".saving"
    try:
        doc.save(tmp)
        os.replace(tmp, path)
    except OSError:
        # 某些文件系统不允许在目标目录新建临时文件（或目标被占用）：退回直接保存。
        try:
            os.remove(tmp)
        except OSError:
            pass
        doc.save(path)

def style_doc(doc: Document):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)

def add_md(doc: Document, md_text: str, base_level: int):
    """Render a light markdown subset: #/##/### headings, tables, fenced code, bullets, bold-stripped text."""
    lines = md_text.splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("```"):
            i += 1
            block = []
            while i < len(lines) and not lines[i].startswith("```"):
                block.append(lines[i]); i += 1
            for b in block:
                p = doc.add_paragraph()
                r = p.add_run(b)
                r.font.name = "Consolas"; r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
                r.font.size = Pt(9)
            i += 1; continue
        if ln.startswith("|") and i + 1 < len(lines) and set(lines[i + 1].replace("|", "").replace("-", "").replace(" ", "")) == set():
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                rows.append(cells); i += 1
            rows.pop(1)  # separator
            if rows:
                t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Table Grid"
                for ri, row in enumerate(rows):
                    for ci, val in enumerate(row[:len(t.columns)]):
                        cell = t.cell(ri, ci); cell.text = val.replace("**", "")
                        if ri == 0:
                            for r_ in cell.paragraphs[0].runs: r_.bold = True
                doc.add_paragraph("")
            continue
        if ln.startswith("#"):
            level = len(ln) - len(ln.lstrip("#"))
            text = ln.lstrip("#").strip()
            doc.add_heading(text, level=min(base_level + level - 1, 4))
        elif ln.strip():
            text = ln.strip()
            style = "List Bullet" if text.startswith(("- ", "* ")) else None
            text = text.lstrip("- ").lstrip("* ")
            text = text.replace("**", "").replace("`", "")
            p = doc.add_paragraph(text, style=style)
        i += 1

def strip_default_bibliography_style(path) -> None:
    """python-docx 默认写入的 customXml 会把 Word 的"引文样式"标成 APA.XSL。
    本工具的文档与参考文献样式无关，这里把该默认值清空，避免文档里残留无关的样式名。"""
    import os, shutil, zipfile
    path = str(path)
    tmp = path + ".tmpdir"
    shutil.rmtree(tmp, ignore_errors=True)
    os.makedirs(tmp)
    try:
        with zipfile.ZipFile(path) as z:
            names = z.namelist()
            z.extractall(tmp)
        changed = False
        for root, _, files in os.walk(tmp):
            for fn in files:
                if not fn.endswith(".xml"):
                    continue
                fp = os.path.join(root, fn)
                try:
                    text = open(fp, encoding="utf-8").read()
                except Exception:
                    continue
                if "APA" in text:
                    new = text.replace('SelectedStyle="/APA.XSL" StyleName="APA"',
                                       'SelectedStyle="" StyleName=""')
                    if new != text:
                        open(fp, "w", encoding="utf-8").write(new)
                        changed = True
        if changed:
            out = path + ".new"
            with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
                for n in names:
                    fp = os.path.join(tmp, n)
                    if os.path.isfile(fp):
                        z.write(fp, n)
            os.replace(out, path)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

def main():
    doc = Document(); style_doc(doc)
    title = doc.add_heading("Psychostat 使用说明（IRT · CTT · 心理统计）", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("统一入口｜三个任务共用一个启动器与 outputs 输出约定", style="Subtitle").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_heading("总入口与共同说明", 1)
    doc.add_paragraph(
        "在 Psychostat 文件夹打开 PowerShell 运行 .\\run_psychostat.ps1，先选择任务：1. 心理统计（平时的统计课内容："
        "t检验/方差分析/相关回归/卡方/非参数/中介调节，含SPSS对照教学与中英文结果报告）；2. CTT（经典测量理论：量表项目分析、"
        "信效度、EFA、CFA）；3. IRT（项目反应理论：Rasch/2PL/3PL/GRM/MIRT）。也可以直接运行对应启动器："
        "run_stats_analysis.ps1 / run_ctt_analysis.ps1 / run_irt_analysis.ps1。")
    doc.add_paragraph(
        "共同要求：安装 R（启动器会自动在 D 盘、PATH、注册表与 C 盘依次查找；未装时首次运行会询问是否自动下载安装，包会自动装好）；Word 结果报告需要 Python（含 pandas 与 python-docx，未装时首次运行会询问是否自动安装）。"
        "所有结果保存在 outputs\\<任务>_<标签>_<时间戳>\\，不会覆盖以往分析；每个结果目录含配置快照与运行日志。")

    doc.add_heading("第一部分 心理统计（SPSS 对照教学）", 1)
    add_md(doc, (ROOT / "docs" / "心理统计使用说明.md").read_text(encoding="utf-8"), base_level=2)
    doc.add_page_break()

    doc.add_heading("第二部分 CTT（经典测量理论）", 1)
    add_md(doc, (ROOT / "docs" / "CTT使用说明.md").read_text(encoding="utf-8"), base_level=2)
    doc.add_page_break()

    doc.add_heading("第三部分 IRT（项目反应理论）", 1)
    add_md(doc, (ROOT / "docs" / "IRT工具使用说明.md").read_text(encoding="utf-8"), base_level=2)

    out = ROOT / "Psychostat使用说明.docx"
    save_docx(doc, out)
    strip_default_bibliography_style(out)
    print("MERGED_MANUAL_OK", out)

if __name__ == "__main__":
    main()

