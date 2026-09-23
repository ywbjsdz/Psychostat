"""Create bilingual results-style Word reports from a Generic IRT result folder.

The script is deliberately defensive: optional diagnostics are included when they
exist, and a clear exception is printed by the PowerShell launcher if generation
cannot proceed.
"""
import argparse
import base64
import sys
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


def border(cell, **edges):
    """Apply results-style top/header/bottom horizontal rules to a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge, value in edges.items():
        element = borders.find(qn("w:" + edge))
        if element is None:
            element = OxmlElement("w:" + edge)
            borders.append(element)
        element.set(qn("w:val"), value)
        element.set(qn("w:sz"), "8")
        element.set(qn("w:color"), "000000")


def setup(doc):
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 2


def add_table(doc, df, number, title, note=""):
    """Write a compact three-line table, or a clear no-data notice."""
    doc.add_paragraph().add_run(f"Table {number}").bold = True
    doc.add_paragraph().add_run(title).italic = True
    if df is None or df.empty or len(df.columns) == 0:
        paragraph(doc, "No data were available for this table." if title.isascii() else "本表没有可用数据。")
        return
    out = df.copy()
    for column in out.columns:
        if pd.api.types.is_numeric_dtype(out[column]):
            out[column] = out[column].map(lambda x: "" if pd.isna(x) else f"{x:.3f}")
    table = doc.add_table(rows=1, cols=len(out.columns))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, column in enumerate(out.columns):
        cell = table.rows[0].cells[index]
        cell.text = str(column)
        border(cell, top="single", bottom="single")
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(8)
    for row_number, (_, row) in enumerate(out.iterrows()):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = str(value)
            for run in cells[index].paragraphs[0].runs:
                run.font.size = Pt(8)
            if row_number == len(out) - 1:
                border(cells[index], bottom="single")
    if note:
        note_paragraph = doc.add_paragraph()
        note_paragraph.add_run("Note. " + note).italic = True


def paragraph(doc, text):
    item = doc.add_paragraph(text)
    item.paragraph_format.first_line_indent = Inches(0.5)
    return item


def figure(doc, path, caption):
    if path.exists() and path.stat().st_size > 0:
        image = doc.add_paragraph()
        image.alignment = WD_ALIGN_PARAGRAPH.CENTER
        image.add_run().add_picture(str(path), width=Inches(6.2))
        label = doc.add_paragraph(caption)
        label.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if label.runs:
            label.runs[0].italic = True


def read_required(result, prefix, suffix):
    path = result / f"{prefix}_{suffix}.csv"
    if not path.exists():
        raise FileNotFoundError(f"Required analysis output is missing: {path}")
    return pd.read_csv(path)


def read_optional(result, prefix, suffix):
    path = result / f"{prefix}_{suffix}.csv"
    return pd.read_csv(path) if path.exists() else None


def build(result, zh):
    matches = list(result.glob("*_model_comparison.csv"))
    if len(matches) != 1:
        raise RuntimeError(
            f"Expected exactly one '*_model_comparison.csv' in {result}; found {len(matches)}."
        )
    prefix = matches[0].name.removesuffix("_model_comparison.csv")
    models = read_required(result, prefix, "model_comparison")
    overview = read_required(result, prefix, "item_overview")
    params = read_required(result, prefix, "item_parameters")
    theta = read_required(result, prefix, "ability_summary")
    missing = read_optional(result, prefix, "missingness_summary")
    item_fit = read_optional(result, prefix, "item_fit")
    reliability = read_optional(result, prefix, "reliability")
    correlations = read_optional(result, prefix, "dimension_correlations")
    eifa_comparison = read_optional(result, prefix, "eifa_dimension_comparison")
    eifa_loadings = read_optional(result, prefix, "eifa_rotated_loadings")
    eifa_eigenvalues = read_optional(result, prefix, "eifa_eigenvalues")
    cifa_loadings = read_optional(result, prefix, "cifa_loading_matrix")
    cifa_significance = read_optional(result, prefix, "cifa_loading_significance")
    cifa_comparison = read_optional(result, prefix, "cifa_vs_unidimensional")

    document = Document()
    setup(document)
    selected = str(models.iloc[0]["Model"]) if not models.empty else "not available"

    if zh:
        title = "通用 IRT 自动分析报告"
        abstract = (
            f"本报告比较了配置的 IRT 候选模型，并按预设准则选择 {selected}。"
            "报告呈现数据与缺失值描述、模型拟合、项目参数和能力估计。"
            "实质性解释仍应结合量表理论、题目内容和研究设计。"
        )
        method = "数据来源、项目列、模型、估计器、迭代上限与输出设置均保存在 config_snapshot.yaml 中，以支持可复现性。"
        missing_title = "数据与缺失值概览"
        missing_note = (
            "缺失单元格数与缺失比例（Missing_cells_before / Missing_pct_before）在剔除整行全缺失被试之前统计；"
            "All_missing_removed 为剔除的整行全缺失被试数。pairwise 与 none 对模型拟合等价"
            "（均基于观测反应模式、不做插补），pairwise 仅影响维度启发式中的成对相关（pairwise.complete.obs）。"
        )
        overview_title = "项目反应概览"
        overview_note = "Missing_pct 为进入模型分析样本中各项目的缺失百分比。"
        results = "结果"
        models_title = "候选模型比较与整体拟合"
        models_note = "AIC、BIC 与 SABIC 越低表示相对拟合越好。"
        params_title = "项目参数"
        params_note = "阈值参数位于模型潜变量量尺上。"
        fit_title = "项目拟合诊断"
        fit_note = "BH 校正后的 p 值用于提示后续核查，而非自动删除项目。"
        theta_title = "能力估计描述统计"
        theta_note = "MAP 能力估计位于拟合模型的潜变量量尺上；N_valid 为该维度上有效（非缺失）估计的人数。"
        reliability_title = "信度估计"
        reliability_note = (
            "边际信度（marginal_rxx）基于模型隐含的误差分布，仅适用于单维模型；"
            "经验信度（empirical_rxx）基于 MAP 能力估计值及其标准误逐维度计算；NA 表示当前模型/估计方式下不可用。"
        )
        corr_title = "潜变量维度相关矩阵"
        corr_note = "对角线为 1.00。"
        discussion = (
            "模型选择反映当前数据与预设比较准则，不能替代理论驱动的结构评估。"
            "建议进一步评估反向计分、局部独立性、DIF、外部效标及跨样本稳定性。"
            "能力估计采用 MAP（最大后验）方法，存在向均值收缩（回归均值）的倾向，极端能力会被低估；解读极端组时应结合标准误（SE）。"
            "表格采用 结果报告风格三线表，图形以 300 dpi PNG 保存。"
        )
        captions = [
            "图 1. 项目反应分布。",
            "图 2. 项目特征曲线（MIRT 为条件切片）。",
            "图 3. 项目信息函数（MIRT 为条件切片）。",
            "图 4. 测验信息函数（MIRT 为条件切片）。",
            "图 5. MAP 能力估计分布。",
        ]
    else:
        title = "Generic Automated IRT Analysis Report"
        abstract = (
            f"This report compared the configured IRT candidate models and selected {selected} using the prespecified criterion. "
            "It presents data and missingness description, model fit, item parameters, and ability estimates. "
            "Substantive interpretation should be integrated with scale theory, item content, and study design."
        )
        method = "The data source, item columns, models, estimator, iteration limit, and output settings are preserved in config_snapshot.yaml to support reproducibility."
        missing_title = "Data and missingness overview"
        missing_note = (
            "Missing cells and percentages (Missing_cells_before / Missing_pct_before) are computed before excluding "
            "respondents with no observed response; All_missing_removed counts those excluded respondents. "
            "Pairwise is equivalent to none for model fitting (observed response patterns, no imputation) and only affects "
            "the pairwise.complete.obs correlations used by the dimensionality heuristic."
        )
        overview_title = "Item-response overview"
        overview_note = "Missing_pct is the item-level missing-response percentage in the analysed sample."
        results = "Results"
        models_title = "Candidate-model comparison and global fit"
        models_note = "Lower AIC, BIC, and SABIC indicate better relative fit."
        params_title = "Item parameters"
        params_note = "Threshold parameters are on the fitted latent metric."
        fit_title = "Item-fit diagnostics"
        fit_note = "BH-adjusted p values flag items for follow-up, not automatic deletion."
        theta_title = "Ability-estimate descriptives"
        theta_note = "MAP ability estimates are on the fitted latent metric; N_valid is the number of valid (non-missing) estimates per dimension."
        reliability_title = "Reliability estimates"
        reliability_note = (
            "Marginal reliability (marginal_rxx) is derived from the model-implied error distribution and is defined for unidimensional models only; "
            "empirical reliability (empirical_rxx) is computed per dimension from the MAP ability estimates and their standard errors; "
            "NA indicates the value is unavailable for the selected model or scoring method."
        )
        corr_title = "Latent-dimension correlation matrix"
        corr_note = "Diagonal values equal 1.00."
        discussion = (
            "Model selection reflects the current data and prespecified comparison criterion; it does not replace theory-driven structural evaluation. "
            "Further work should assess reverse scoring, local dependence, DIF, external criteria, and cross-sample stability. "
            "MAP ability estimates shrink toward the mean, so extreme abilities are underestimated; interpret extreme groups alongside their standard errors. "
            "Tables use an results-style three-line layout and figures are saved as 300 dpi PNG files."
        )
        captions = [
            "Figure 1. Item-response distributions.",
            "Figure 2. Item characteristic curves (conditional slices for MIRT).",
            "Figure 3. Item information functions (conditional slices for MIRT).",
            "Figure 4. Test information functions (conditional slices for MIRT).",
            "Figure 5. Distribution of MAP ability estimates.",
        ]

    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.add_run(title).bold = True
    heading.runs[0].font.size = Pt(14)
    document.add_heading("摘要" if zh else "Abstract", level=1)
    paragraph(document, abstract)
    document.add_heading("方法" if zh else "Method", level=1)
    paragraph(document, method)
    table_number = 1
    if missing is not None:
        add_table(document, missing, table_number, missing_title, missing_note)
        table_number += 1
    add_table(document, overview, table_number, overview_title, overview_note)
    table_number += 1
    document.add_heading(results, level=1)
    add_table(document, models, table_number, models_title, models_note)
    table_number += 1
    if reliability is not None:
        add_table(document, reliability, table_number, reliability_title, reliability_note)
        table_number += 1
    add_table(document, params, table_number, params_title, params_note)
    table_number += 1
    if eifa_comparison is not None:
        add_table(document, eifa_comparison, table_number, "探索性维度比较（EIFA）" if zh else "Exploratory dimensionality comparison (EIFA)", "按 BIC 排序；较低值表示更优的相对拟合。" if zh else "Ranked by BIC; lower values indicate better relative fit.")
        table_number += 1
    if eifa_loadings is not None:
        add_table(document, eifa_loadings, table_number, "旋转后因子载荷（EIFA）" if zh else "Rotated factor loadings (EIFA)", "采用 oblimin 斜交旋转；h2 为共同度。" if zh else "Oblimin oblique rotation; h2 is communality.")
        table_number += 1
    if eifa_eigenvalues is not None:
        add_table(document, eifa_eigenvalues, table_number, "特征值（EIFA 诊断）" if zh else "Eigenvalues (EIFA diagnostic)", "基于项目反应相关矩阵，仅作为维度探索的辅助证据。" if zh else "Based on the item-response correlation matrix; auxiliary evidence only.")
        table_number += 1
    if cifa_loadings is not None:
        add_table(document, cifa_loadings, table_number, "约束后因子载荷（CIFA）" if zh else "Constrained factor loadings (CIFA)", "未指定的交叉载荷在模型中固定为零。" if zh else "Unspecified cross-loadings are fixed to zero in the model.")
        table_number += 1
    if cifa_significance is not None:
        add_table(document, cifa_significance, table_number, "载荷显著性检验（CIFA）" if zh else "Loading significance tests (CIFA)", "z = 载荷估计值 / 标准误；p 值为双侧 Wald 近似。" if zh else "z = loading estimate / SE; p values are two-sided Wald approximations.")
        table_number += 1
    if cifa_comparison is not None:
        add_table(document, cifa_comparison, table_number, "CIFA 与单维模型比较" if zh else "CIFA versus unidimensional model", "Δ 值定义为 CIFA − 单维模型；CFI 增加、RMSEA/AIC/BIC 减少通常支持多维结构。" if zh else "Deltas are CIFA minus the unidimensional model; increased CFI and reduced RMSEA/AIC/BIC generally support multidimensional structure.")
        table_number += 1
    if item_fit is not None:
        add_table(document, item_fit, table_number, fit_title, fit_note)
        table_number += 1
    add_table(document, theta, table_number, theta_title, theta_note)
    table_number += 1
    if correlations is not None:
        add_table(document, correlations, table_number, corr_title, corr_note)

    for suffix, caption in zip(
        ["response_distributions", "icc", "iif", "tif", "theta_distribution"], captions
    ):
        figure(document, result / f"{prefix}_{suffix}.png", caption)

    document.add_heading("讨论与限制" if zh else "Discussion and limitations", level=1)
    paragraph(document, discussion)
    output = result / (f"{prefix}_report_zh.docx" if zh else f"{prefix}_report_en.docx")
    document.save(output)
    print(f"Created report: {output}")


if __name__ == "__main__":
    # Output as UTF-8 where the host console supports it; paths themselves are
    # passed as ASCII-only Base64, so they do not depend on a console code page.
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="backslashreplace")
    if hasattr(sys.stderr, "reconfigure"):
        sys.stderr.reconfigure(encoding="utf-8", errors="backslashreplace")

    parser = argparse.ArgumentParser()
    selector = parser.add_mutually_exclusive_group(required=True)
    selector.add_argument("--result-dir", help="Legacy direct path argument.")
    selector.add_argument(
        "--result-dir-utf8-base64",
        help="UTF-8 result path encoded as Base64; safe for Windows console code pages.",
    )
    arguments = parser.parse_args()
    if arguments.result_dir_utf8_base64:
        try:
            result_text = base64.b64decode(arguments.result_dir_utf8_base64, validate=True).decode("utf-8")
        except (ValueError, UnicodeDecodeError) as error:
            raise ValueError("Invalid UTF-8 Base64 result-directory argument.") from error
    else:
        result_text = arguments.result_dir
    result_dir = Path(result_text)
    if not result_dir.is_dir():
        raise NotADirectoryError(f"Result directory does not exist: {result_dir}")
    build(result_dir, True)
    build(result_dir, False)