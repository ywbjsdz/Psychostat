"""把 docs/CTT与IRT校对记录.md 渲染为项目根目录的 CTT与IRT校对记录.docx。

复用 scripts/merge_manuals.py 里既有的 markdown→docx 渲染器（标题/表格/列表/代码），
保证与三份使用说明书的排版口径一致，不引入新的依赖。
用法（Windows）：py -3 scripts/make_verification_record_docx.py
"""
from __future__ import annotations
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))

from docx import Document                                   # noqa: E402
from merge_manuals import add_md, style_doc, save_docx       # noqa: E402


def main() -> int:
    src = ROOT / "docs" / "CTT与IRT校对记录.md"
    if not src.exists():
        print(f"找不到源文件：{src}")
        return 1
    doc = Document()
    style_doc(doc)
    add_md(doc, src.read_text(encoding="utf-8"), base_level=1)
    out = ROOT / "CTT与IRT校对记录.docx"
    save_docx(doc, out)
    print(f"已生成：{out}（{out.stat().st_size} 字节，段落 {len(doc.paragraphs)} 个，表格 {len(doc.tables)} 个）")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
