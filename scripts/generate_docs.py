r"""Render Psychostat's Markdown guides into Word documents.

用法：
    python scripts\generate_docs.py            # 渲染全部
    python scripts\generate_docs.py manual     # 只渲染说明书

输出（Word 版：说明书放项目根目录便于双击，其余文档放 docs/；Markdown 源文件保留原位便于 diff）：
    docs\Psychostat_完整使用说明书.md  →  Psychostat完整使用说明书.docx
    （待修问题清单只保留 md 源；其 docx 属可再生产物，不随仓库发布）

实现复用 merge_manuals.py 的轻量 Markdown 渲染器（标题/表格/代码块/列表），保证排版风格与
既有的 Psychostat使用说明.docx 一致；并在渲染前把块引用与锚点目录降级为普通文本，避免 Word 里
出现 ">" 与 "[标题](#锚点)" 这类原始记号。
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(Path(__file__).resolve().parent))
from merge_manuals import add_md, style_doc  # noqa: E402

LINK_RE = re.compile(r"\[([^\]]+)\]\(#[^)]*\)")

DOCS = [
    {
        "key": "manual",
        "src": ROOT / "docs" / "Psychostat_完整使用说明书.md",
        "dst": ROOT / "Psychostat完整使用说明书.docx",
        "title": "Psychostat 完整使用说明书",
        "subtitle": ["心理统计 · CTT 经典测量理论 · IRT 项目反应理论",
                     "面向刚学心理统计的同学与从事量表/测验研究的研究者"],
    },
]



def save_docx(doc, path) -> None:
    """先写到同目录临时文件再原子替换：避免在部分文件系统上直接原地保存失败。"""
    import os
    path = str(path)
    tmp = path + ".saving"
    try:
        doc.save(tmp)
        os.replace(tmp, path)
    except OSError:
        # 某些文件系统不允许在目标目录新建临时文件（或目标被占用）：退回直接保存。
        try:
            os.remove(tmp)
        except OSError:
            pass
        doc.save(path)

def renderable(text: str) -> str:
    """Drop block-quote markers and anchor links so the light renderer stays clean."""
    out = []
    for line in text.splitlines():
        if line.startswith(">"):
            line = line.lstrip(">").lstrip()
        out.append(LINK_RE.sub(r"\1", line))
    return "\n".join(out)


def build(doc_spec: dict) -> Path | None:
    src, dst = doc_spec["src"], doc_spec["dst"]
    if not src.exists():
        # 源 md 未随工程发布（例如开发用的问题清单已被移除）时跳过，不影响其余文档生成
        print(f"DOC_SKIP {src} (source not found)")
        return None
    doc = Document()
    style_doc(doc)
    doc.add_heading(doc_spec["title"], 0).alignment = WD_ALIGN_PARAGRAPH.CENTER
    for line in doc_spec["subtitle"]:
        doc.add_paragraph(line).alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_md(doc, renderable(src.read_text(encoding="utf-8")), base_level=1)
    save_docx(doc, dst)
    strip_default_bibliography_style(dst)
    return dst


def strip_default_bibliography_style(path) -> None:
    """python-docx 默认写入的 customXml 会把 Word 的"引文样式"标成 APA.XSL。
    本工具的文档与参考文献样式无关，这里把该默认值清空，避免文档里残留无关的样式名。"""
    import os, shutil, zipfile
    path = str(path)
    tmp = path + ".tmpdir"
    shutil.rmtree(tmp, ignore_errors=True)
    os.makedirs(tmp)
    try:
        with zipfile.ZipFile(path) as z:
            names = z.namelist()
            z.extractall(tmp)
        changed = False
        for root, _, files in os.walk(tmp):
            for fn in files:
                if not fn.endswith(".xml"):
                    continue
                fp = os.path.join(root, fn)
                try:
                    text = open(fp, encoding="utf-8").read()
                except Exception:
                    continue
                if "APA" in text:
                    new = text.replace('SelectedStyle="/APA.XSL" StyleName="APA"',
                                       'SelectedStyle="" StyleName=""')
                    if new != text:
                        open(fp, "w", encoding="utf-8").write(new)
                        changed = True
        if changed:
            out = path + ".new"
            with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
                for n in names:
                    fp = os.path.join(tmp, n)
                    if os.path.isfile(fp):
                        z.write(fp, n)
            os.replace(out, path)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def main() -> None:
    wanted = [a.lower() for a in sys.argv[1:]]
    built = [p for p in (build(d) for d in DOCS if not wanted or d["key"] in wanted) if p]
    for path in built:
        print("DOC_OK", path)
    if not built:
        raise SystemExit(f"nothing built; valid keys: {[d['key'] for d in DOCS]}")


if __name__ == "__main__":
    main()
